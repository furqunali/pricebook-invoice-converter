"""Render the Phase-5 reports to disk: Excel, PDF, and a Power BI dataset.

Kept separate from ``reporting.py`` (the data engine) so the parsing/aggregation
stays import-light; openpyxl and reportlab are imported lazily here (the
``reports`` extra). Everything is subscription/offline — no API.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional

from . import reporting as R
from .config import Settings


def _stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #

def _xl_style_header(ws, row, ncols):
    from openpyxl.styles import Alignment, Font, PatternFill
    fill = PatternFill("solid", fgColor=R.NAVY)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")


def write_excel(agg: R.Aggregates, recon_rows, recon_status, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # -- Summary ------------------------------------------------------------ #
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Demo Retail Co"
    ws["A1"].font = Font(size=14, bold=True, color="000000")
    ws["A2"] = f"Period: {agg.period}"
    ws["A3"] = f"Generated: {_stamp()}"
    kpis = [
        ("Total purchases", f"${agg.total_purchases():,.2f}"),
        ("Invoices", len(agg.invoices)),
        ("Sites", len(agg.by_site())),
        ("Vendors", len(agg.by_vendor())),
        ("Review rate", f"{agg.review_rate()*100:.1f}%"),
    ]
    ws["A5"] = "KPI"; ws["B5"] = "Value"
    _xl_style_header(ws, 5, 2)
    for i, (k, v) in enumerate(kpis, start=6):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.freeze_panes = "A6"

    # small data block for charts (by site)
    sites = agg.by_site()
    base = 13
    ws.cell(row=base, column=1, value="Site")
    ws.cell(row=base, column=2, value="Purchases")
    _xl_style_header(ws, base, 2)
    for i, s in enumerate(sites, start=base + 1):
        ws.cell(row=i, column=1, value=s["site"])
        ws.cell(row=i, column=2, value=R._f(s["total"])).number_format = R._MONEY
    if sites:
        pie = PieChart(); pie.title = "Purchases by store"
        data = Reference(ws, min_col=2, min_row=base, max_row=base + len(sites))
        cats = Reference(ws, min_col=1, min_row=base + 1, max_row=base + len(sites))
        pie.add_data(data, titles_from_data=True); pie.set_categories(cats)
        pie.height, pie.width = 7, 11
        ws.add_chart(pie, "D5")

    # -- Store-wise --------------------------------------------------------- #
    _sheet_table(wb, "Store-wise",
                 ["Site ID", "Site", "Invoices", "Purchases", "Returns (credit)"],
                 [[s["site_id"], s["site"], s["count"], R._f(s["total"]), R._f(s["returns"])]
                  for s in sites],
                 money_cols=[4, 5], total_cols=[3, 4, 5])

    # -- Vendor-wise (+ bar chart of top 10) -------------------------------- #
    vendors = agg.by_vendor()
    vw = _sheet_table(wb, "Vendor-wise",
                      ["Vendor ID", "Vendor", "Invoices", "Purchases"],
                      [[v["vendor_id"], v["vendor"], v["count"], R._f(v["total"])]
                       for v in vendors],
                      money_cols=[4], total_cols=[3, 4])
    topn = min(10, len(vendors))
    if topn:
        bar = BarChart(); bar.title = f"Top {topn} vendors by purchases"; bar.legend = None
        data = Reference(vw, min_col=4, min_row=1, max_row=1 + topn)
        cats = Reference(vw, min_col=2, min_row=2, max_row=1 + topn)
        bar.add_data(data, titles_from_data=True); bar.set_categories(cats)
        bar.height, bar.width = 8, 16
        vw.add_chart(bar, f"F2")

    # -- Reconciliation ----------------------------------------------------- #
    rec = _sheet_table(wb, "Reconciliation",
                       ["Vendor", "Converter", "PDI Posted", "Vendor Stmt", "Variance"],
                       [[r["vendor"], R._f(r["converter"]), R._money(r["pdi_posted"]),
                         R._money(r["vendor_stmt"]),
                         (R._money(r["variance"]) if r["variance"] is not None else "PENDING")]
                        for r in recon_rows],
                       money_cols=[2], total_cols=[2])
    rec.cell(row=rec.max_row + 2, column=1,
             value=f"3-way match status: {recon_status}. PDI-Export and vendor "
                   f"statements are configurable stubs (pdi_export_dir / "
                   f"vendor_statements_dir) for Part 2.")

    # -- Exceptions --------------------------------------------------------- #
    _sheet_table(wb, "Exceptions", ["Source file", "Reason"],
                 [[src, reason] for src, reason in agg.exceptions()] or [["(none)", "all clear"]])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _sheet_table(wb, title, headers, rows, money_cols=(), total_cols=()):
    from openpyxl.styles import Font
    ws = wb.create_sheet(title)
    ws.append(headers)
    _xl_style_header(ws, 1, len(headers))
    for row in rows:
        ws.append(row)
    for c in money_cols:
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=c).number_format = R._MONEY
    if rows and total_cols:
        tr = ws.max_row + 1
        ws.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True)
        for c in total_cols:
            vals = [ws.cell(row=r, column=c).value for r in range(2, tr)]
            nums = [v for v in vals if isinstance(v, (int, float))]
            cell = ws.cell(row=tr, column=c, value=sum(nums))
            cell.font = Font(bold=True)
            if c in money_cols:
                cell.number_format = R._MONEY
    for i, h in enumerate(headers, start=1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(str(h)) + 4)
    ws.freeze_panes = "A2"
    return ws


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

def write_pdf(agg: R.Aggregates, recon_rows, recon_status, out: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.piecharts import Pie

    navy = colors.HexColor("#" + R.NAVY)
    teal = colors.HexColor("#" + R.TEAL)
    light = colors.HexColor("#" + R.LIGHT)

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], textColor=colors.black, fontSize=18)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=teal)
    small = ParagraphStyle("small", parent=styles["Normal"], textColor=colors.grey, fontSize=8)

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.grey); canvas.setFont("Helvetica", 8)
        canvas.drawString(0.75 * inch, 0.5 * inch,
                          "Demo Retail Co Invoice Converter — confidential")
        canvas.drawRightString(7.75 * inch, 0.5 * inch, f"Page {doc.page}")
        canvas.restoreState()

    def _tbl(headers, rows, money_idx=(), widths=None):
        data = [headers] + rows
        t = Table(data, colWidths=widths, repeatRows=1)
        st = [
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        t.setStyle(TableStyle(st))
        return t

    story = []
    story.append(Paragraph("Demo Retail Co", h1))
    story.append(Paragraph(f"Period: <b>{agg.period}</b> &nbsp;|&nbsp; Generated: {_stamp()}", small))
    story.append(Spacer(1, 10))

    # KPI strip
    kpis = [["Total purchases", "Invoices", "Sites", "Vendors", "Review rate"],
            [f"${agg.total_purchases():,.2f}", str(len(agg.invoices)),
             str(len(agg.by_site())), str(len(agg.by_vendor())),
             f"{agg.review_rate()*100:.1f}%"]]
    kt = Table(kpis, colWidths=[1.5 * inch] * 5)
    kt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), teal),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 1), (-1, 1), light),
        ("FONTSIZE", (0, 0), (-1, -1), 9), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.white), ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(kt)
    story.append(Spacer(1, 14))

    # Store-wise + pie
    sites = agg.by_site()
    story.append(Paragraph("Store-wise purchases", h2))
    rows = [[s["site_id"], s["site"], str(s["count"]), f"${s['total']:,.2f}",
             f"${s['returns']:,.2f}"] for s in sites]
    rows.append(["", "TOTAL", str(len(agg.invoices)),
                 f"${agg.total_purchases():,.2f}", ""])
    story.append(_tbl(["Site ID", "Site", "Invoices", "Purchases", "Returns"], rows,
                      widths=[0.8 * inch, 2.2 * inch, 0.9 * inch, 1.3 * inch, 1.1 * inch]))
    if sites:
        d = Drawing(360, 170)
        pie = Pie(); pie.x, pie.y = 120, 10; pie.width = pie.height = 150
        pie.data = [R._f(s["total"]) for s in sites]
        pie.labels = [s["site"] for s in sites]
        pie.slices.strokeWidth = 0.5
        palette = ["#0E7C7B", "#1F3A5F", "#3AA6A4", "#F2A65A", "#7B6CE0",
                   "#E06C9F", "#5FB0E0", "#9BC53D"]
        for i in range(len(sites)):
            pie.slices[i].fillColor = colors.HexColor(palette[i % len(palette)])
        d.add(pie)
        story.append(Spacer(1, 8)); story.append(d)
    story.append(Spacer(1, 12))

    # Top vendors
    vendors = agg.by_vendor()[:10]
    story.append(Paragraph("Top vendors by purchases", h2))
    vrows = [[v["vendor_id"], v["vendor"], str(v["count"]), f"${v['total']:,.2f}"]
             for v in vendors]
    story.append(_tbl(["Vendor ID", "Vendor", "Invoices", "Purchases"], vrows,
                      widths=[0.9 * inch, 3.0 * inch, 1.0 * inch, 1.4 * inch]))
    story.append(Spacer(1, 12))

    # Reconciliation framework
    story.append(Paragraph("Reconciliation framework (3-way)", h2))
    story.append(Paragraph(f"Status: <b>{recon_status}</b> — PDI-Export and vendor "
                           "statements are configurable stubs for Part 2.", small))
    rr = [[r["vendor"], f"${r['converter']:,.2f}", R._money(r["pdi_posted"]),
           R._money(r["vendor_stmt"]),
           (R._money(r["variance"]) if r["variance"] is not None else "PENDING")]
          for r in recon_rows[:15]]
    story.append(_tbl(["Vendor", "Converter", "PDI Posted", "Vendor Stmt", "Variance"], rr,
                      widths=[2.4 * inch, 1.2 * inch, 1.2 * inch, 1.2 * inch, 1.1 * inch]))
    story.append(Spacer(1, 12))

    # Exceptions
    story.append(Paragraph("Exceptions (still in review)", h2))
    exc = agg.exceptions()
    if exc:
        erows = [[src, (reason[:80] + "…") if len(reason) > 80 else reason] for src, reason in exc]
        story.append(_tbl(["Source file", "Reason"], erows,
                          widths=[2.2 * inch, 5.0 * inch]))
    else:
        story.append(Paragraph("All clear — nothing in review.", styles["Normal"]))

    out.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(out), pagesize=letter,
                            topMargin=0.7 * inch, bottomMargin=0.8 * inch)
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return out


# --------------------------------------------------------------------------- #
# Power BI dataset (refreshable flat CSV + suggested DAX)
# --------------------------------------------------------------------------- #

def write_powerbi(agg: R.Aggregates, out_dir: Path) -> list[Path]:
    import csv as _csv
    out_dir.mkdir(parents=True, exist_ok=True)
    fact = out_dir / "purchase_fact.csv"
    with open(fact, "w", newline="", encoding="utf-8") as f:
        w = _csv.writer(f)
        w.writerow(["date", "batch_id", "vendor_id", "vendor", "site_id", "site",
                    "invoice_ref", "net_total", "gross", "returns", "discounts", "line_count"])
        for iv in agg.invoices:
            w.writerow([iv.date, iv.batch_id, iv.vendor_id,
                        agg.vendor_names.get(iv.vendor_id, iv.vendor_id),
                        iv.site_id, agg.site_names.get(iv.site_id, iv.site_id),
                        iv.invoice_ref, f"{iv.total:.2f}", f"{iv.gross:.2f}",
                        f"{iv.returns:.2f}", f"{iv.discounts:.2f}", len(iv.lines)])
    dax = out_dir / "measures_dax.md"
    dax.write_text(
        "# Power BI — purchase_fact dataset\n\n"
        "Refresh: point Power BI at `purchase_fact.csv` (Get Data > Text/CSV) or the\n"
        "folder; the reconciliation job rewrites it each run, so **Refresh** re-reads it.\n\n"
        "## Suggested measures (DAX)\n"
        "```DAX\n"
        "Total Purchases = SUM(purchase_fact[net_total])\n"
        "Total Returns   = SUM(purchase_fact[returns])\n"
        "Total Discounts = SUM(purchase_fact[discounts])\n"
        "Invoice Count   = DISTINCTCOUNT(purchase_fact[invoice_ref])\n"
        "Avg Invoice     = DIVIDE([Total Purchases], [Invoice Count])\n"
        "```\n\n"
        "## Suggested visuals\n"
        "- Purchases by **site** (donut), by **vendor** (bar, top N)\n"
        "- Purchases **trend** by date; Returns & Discounts by vendor\n"
        "- Converted-vs-posted once the PDI-Export feed is wired (Part 2)\n",
        encoding="utf-8")
    return [fact, dax]


# --------------------------------------------------------------------------- #
# Orchestrator
# --------------------------------------------------------------------------- #

def build_reports(settings: Optional[Settings] = None,
                  converted_dir: Optional[Path] = None,
                  logs_dir: Optional[Path] = None,
                  out_dir: Optional[Path] = None,
                  period: Optional[str] = None) -> dict:
    """Build PDF + Excel + Power BI dataset from the converter's output. Returns
    the paths written. No API — subscription/offline only."""
    from .paths import Paths
    settings = settings or Settings.load()
    paths = Paths.from_settings(settings)
    converted_dir = converted_dir or paths.converted
    logs_dir = logs_dir or paths.logs
    out_dir = out_dir or paths.reports
    period = period or datetime.now().strftime("%Y-%m")

    agg = R.build_aggregates(converted_dir, logs_dir, settings, period)
    recon_rows, recon_status = R.reconciliation_rows(agg, settings)

    safe = period.replace("/", "-").replace(" ", "_")
    xlsx = write_excel(agg, recon_rows, recon_status, out_dir / f"purchase_report_{safe}.xlsx")
    pdf = write_pdf(agg, recon_rows, recon_status, out_dir / f"purchase_report_{safe}.pdf")
    pbi = write_powerbi(agg, out_dir / "powerbi")
    log.info("reports written: %s", [str(xlsx), str(pdf), *map(str, pbi)])
    return {"excel": xlsx, "pdf": pdf, "powerbi": pbi,
            "invoices": len(agg.invoices), "total": agg.total_purchases(),
            "recon_status": recon_status}


log = R.log
